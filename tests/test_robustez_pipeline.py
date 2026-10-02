# -*- coding: utf-8 -*-
"""
Robustez y trazabilidad del pipeline del plan:
- una excepción inesperada se entrega como alerta, no se lanza al llamador;
- el año de corte usado queda en los metadatos (resultado reproducible);
- los objetivos viajan en `to_dict()`;
- el «Enfoque» del bloque IR llega al modelo.
"""

from __future__ import annotations

import os
import sys
import tempfile
from dataclasses import replace
from datetime import datetime

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from extractor_pa import MAPEO_NUEVO, extraer_plan_accion  # noqa: E402
from extractor_pa import pipeline  # noqa: E402
from tests.test_smoke import _construir_excel  # noqa: E402


@pytest.fixture(scope="module")
def plan_sintetico():
    ruta = os.path.join(tempfile.mkdtemp(), "plan_robustez.xlsx")
    _construir_excel(ruta)
    return ruta


def test_excepcion_en_el_motor_queda_como_alerta(plan_sintetico, monkeypatch):
    def revienta(*_a, **_k):
        raise RuntimeError("columna rara")

    monkeypatch.setattr(pipeline._MOTOR, "extraer", revienta)
    res = extraer_plan_accion(plan_sintetico)
    assert not res.exitoso
    [alerta] = res.alertas
    assert alerta.tipo == "error_extraccion" and alerta.nivel == "ERROR"
    assert "RuntimeError: columna rara" in alerta.descripcion


def test_excepcion_en_fichas_no_pierde_el_plan(plan_sintetico, monkeypatch):
    def revienta(_wb):
        raise ValueError("ficha corrupta")

    monkeypatch.setattr(pipeline, "leer_fichas", revienta)
    res = extraer_plan_accion(plan_sintetico)
    assert res.exitoso
    assert len(res.indicadores_resultado) == 2
    opcionales = [a for a in res.alertas if a.tipo == "error_etapa_opcional"]
    assert len(opcionales) == 1 and opcionales[0].nivel == "ADVERTENCIA"
    assert "fichas técnicas" in opcionales[0].descripcion


def test_excepcion_en_reglas_no_pierde_el_plan(plan_sintetico, monkeypatch):
    import extractor_pa.validacion as validacion

    def revienta(*_a, **_k):
        raise KeyError("V7")

    monkeypatch.setattr(validacion, "validar_reglas", revienta)
    res = extraer_plan_accion(plan_sintetico, incluir_reglas_negocio=True)
    assert res.exitoso
    assert any(a.tipo == "error_etapa_opcional" and "reglas de negocio" in a.descripcion
               for a in res.alertas)


def test_anio_corte_explicito_y_por_defecto(plan_sintetico):
    assert extraer_plan_accion(plan_sintetico, anio_vigencia=2024).metadatos.anio_corte == 2024
    assert extraer_plan_accion(plan_sintetico).metadatos.anio_corte == datetime.now().year


def test_anio_corte_reproduce_la_extraccion(plan_sintetico):
    """Volver a pasar `metadatos.anio_corte` da exactamente las mismas metas de vigencia."""
    a = extraer_plan_accion(plan_sintetico)
    b = extraer_plan_accion(plan_sintetico, anio_vigencia=a.metadatos.anio_corte)
    clave = lambda r: [(i.codigo_ir, i.anio_vigencia, i.meta_vigencia_actual,  # noqa: E731
                        i.meta_vigencia_anterior) for i in r.indicadores_resultado]
    assert clave(a) == clave(b)


def test_anio_corte_en_alertas_tempranas(tmp_path):
    res = extraer_plan_accion(tmp_path / "no_existe.xlsx", anio_vigencia=2025)
    assert res.alertas[0].tipo == "apertura"
    assert res.metadatos.anio_corte == 2025


def test_to_dict_incluye_objetivos(plan_sintetico):
    d = extraer_plan_accion(plan_sintetico).to_dict()
    assert [o["codigo"] for o in d["objetivos"]] == ["1", "2"]
    assert d["metadatos"]["anio_corte"] is not None


def test_enfoque_del_ir(tmp_path):
    """Con un ancla «Enfoque» en el bloque IR, el valor llega a `IndicadorResultado.enfoque`
    (así lo trae el formato antiguo; aquí se simula sobre el sintético nuevo)."""
    from openpyxl import load_workbook

    ruta = tmp_path / "plan_enfoque.xlsx"
    _construir_excel(str(ruta))
    wb = load_workbook(ruta)
    ws = wb["Plan de acción"]
    ws.cell(row=11, column=48, value="Enfoque")
    ws.cell(row=12, column=48, value="Poblacional")
    wb.save(ruta)

    mapeo = replace(MAPEO_NUEVO, anclas_repetidas={
        **MAPEO_NUEVO.anclas_repetidas, "Enfoque": ("enfoque_ir", "enfoque_princ")})
    res = extraer_plan_accion(ruta, mapeo=mapeo)
    enfoques = {i.codigo_ir: i.enfoque for i in res.indicadores_resultado}
    assert enfoques == {"1.1": "Poblacional", "2.1": None}
