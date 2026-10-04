# -*- coding: utf-8 -*-
"""Unit tests por etapa (Hito 2): helpers puros de utilidades, vigencia, fichas y validación."""

import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from extractor_pa.utilidades import (a_float, extraer_codigo, es_vigente,
                                     peso_positivo, clave_grupo, tiene_contenido,
                                     elegir_peso_objetivo)
from extractor_pa.lector_filas import prefiltrar_filas
from extractor_pa.vigencia import calcular_vigencia
from extractor_pa.lector_fichas import codigo_de_hoja_ficha
from extractor_pa.pipeline import _es_nombre_politica
from extractor_pa.validacion import _factor, _leer_peso


# ── utilidades.a_float (tolerante, formato europeo) ──
@pytest.mark.parametrize("entrada,esperado", [
    (10, 10.0), (10.5, 10.5),
    ("9%", 9.0), ("$1.234,56", 1234.56),   # europeo: punto miles, coma decimal
    ("74,5", 74.5), ("  12  ", 12.0),
    (None, None), ("n/a", None), ("", None), ("texto", None),
])
def test_a_float(entrada, esperado):
    assert a_float(entrada) == esperado


# ── utilidades.extraer_codigo ──
@pytest.mark.parametrize("texto,niveles,esperado", [
    ("1.1 Aumento de la cobertura", 2, "1.1"),
    ("4.1.5Nombre pegado", 3, "4.1.5"),
    ("5. 1 Productividad", 2, "5.1"),     # separador con espacio → normaliza
    ("2 Objetivo general", 1, "2"),
    ("Sin código", None, None),
    ("3.2.1", None, "3.2.1"),
    ("1 . 1 Separador con espacios", 2, "1.1"),
])
def test_extraer_codigo(texto, niveles, esperado):
    assert extraer_codigo(texto, niveles) == esperado


# Prefijo de tipo P/R/O/OE delante del código (Trata: «P1.1.1Acciones…»;
# Trabajo Decente: «OE1. Promover…»).
@pytest.mark.parametrize("texto,niveles,esperado", [
    ("P1.1.1Acciones encaminadas a promover", 3, "1.1.1"),
    ("P1.1.2. Formación por demanda", 3, "1.1.2"),
    ("P2.1.10. Actividades de capacitación", 3, "2.1.10"),
    ("R1.1 Disminución del número de casos", 2, "1.1"),
    ("r 1.1 en minúscula y con espacio", 2, "1.1"),
    ("P. 3.1.1 con punto y espacio", 3, "3.1.1"),
    ("p.3.1.2", 3, "3.1.2"),
    ("O1 Objetivo con prefijo", 1, "1"),
    ("OE1. Promover principios y derechos", 1, "1"),
    ("oe 2. en minúscula y con espacio", 1, "2"),
    ("OE3.3. Procurar el acceso", 1, "3"),
    ("P1.1.1 Indicador del seguimiento", None, "1.1.1"),   # sin niveles: seguimiento
])
def test_extraer_codigo_tolera_prefijo_de_tipo(texto, niveles, esperado):
    assert extraer_codigo(texto, niveles) == esperado


# La lista de prefijos es cerrada: P, R, O u OE pegados al número.
@pytest.mark.parametrize("texto,niveles", [
    ("Plan 2024 de contingencia", None),
    ("Producto 3 sin código", None),
    ("PR1.1 Dos letras", 2),
    ("OEA1. Tres letras", 1),
    ("Oeste 2024", None),
    ("X1.1.1 Letra fuera de la lista", 3),
    ("P", None),
])
def test_extraer_codigo_no_inventa_codigos(texto, niveles):
    assert extraer_codigo(texto, niveles) is None


@pytest.mark.parametrize("valor,esperado", [
    ("Producto sin código", True), ("TOTAL", True), (0, True), ("4", True),
    (".", False), ("  ", False), ("_", False), ("—", False), (None, False),
])
def test_tiene_contenido(valor, esperado):
    assert tiene_contenido(valor) is esperado


# ── lector_filas.prefiltrar_filas ──
def test_prefiltro_aparta_las_filas_con_texto_sin_codigo():
    filas = [
        (12, ["1.1 Resultado", "1.1.1 Producto"]),
        (13, [None, "P1.1.2 Producto con prefijo"]),
        (14, [None, "Producto sin código"]),
        (15, [".", "."]),
        (16, [None, None, "Total"]),
    ]
    descartadas = []
    conservadas = prefiltrar_filas(filas, 1, 2, descartadas)
    assert [r for r, _ in conservadas] == [12, 13]
    assert [r for r, _ in descartadas] == [14]      # «.» y «Total» fuera de columna: sin aviso
    # Sin la lista, el resultado es el mismo de antes.
    assert prefiltrar_filas(filas, 1, 2) == conservadas


# ── utilidades.es_vigente / peso_positivo / clave_grupo ──
@pytest.mark.parametrize("valor,esperado", [
    ("Vigente", True), ("", True), (None, True),
    ("No Vigente", False), ("No", False),
])
def test_es_vigente(valor, esperado):
    assert es_vigente(valor) is esperado


@pytest.mark.parametrize("valor,esperado", [
    (0.5, True), ("10%", True), (0, False), ("", False), ("abc", False),
])
def test_peso_positivo(valor, esperado):
    assert peso_positivo(valor) is esperado


def test_clave_grupo_usa_codigo():
    assert clave_grupo("1.1 Resultado X") == "1.1"


# ── utilidades.elegir_peso_objetivo: (vigente_ir, peso_ir, peso_objetivo) ──
@pytest.mark.parametrize("candidatos,esperado", [
    # Trabajo Decente OE2: el IR No Vigente que encabeza el objetivo pesa 0.
    ([("No Vigente", 0, 0), ("Vigente", 0.1616, 0.4848)], 0.4848),
    # Trata: el No Vigente trae peso de objetivo 0 y peso de IR vacío.
    ([("No vigente", None, 0), ("Vigente", 0.4, 0.4)], 0.4),
    # Vigente con peso de IR 0 no es la fila autoritativa (como en la ascensión).
    ([("Vigente", 0, 0.1), ("Vigente", 0.3, 0.3)], 0.3),
    # Sin marca de vigencia cuenta como vigente.
    ([(None, 0, 0.1), (None, 0.3, 0.3)], 0.3),
    # Espacio Público OE2: solo el No Vigente trae el peso -> el primero no nulo.
    ([("No vigente", 0, 0.34), ("Vigente", 0.17, None)], 0.34),
    # Ninguno vigente: el primero no nulo (como antes).
    ([("No Vigente", 0, None), ("No Vigente", 0, 0.2), ("No Vigente", 0, 0.5)], 0.2),
    # Un texto no numérico no cuenta como peso.
    ([("Vigente", 0.5, "N/A"), ("Vigente", 0.5, "50%")], "50%"),
    ([("Vigente", 0.5, None)], None),
    ([], None),
])
def test_elegir_peso_objetivo(candidatos, esperado):
    assert elegir_peso_objetivo(candidatos) == esperado


# ── validacion._leer_peso / _factor (escala de la ponderación) ──
@pytest.mark.parametrize("valor,esperado", [
    (0.0377, (0.0377, False)), (25, (25.0, False)), (0, (0.0, False)),
    ("0.0377", (0.0377, False)),             # texto sin «%»: tan ambiguo como el número
    ("2.86%", (2.86, True)), ("0,5%", (0.5, True)), (" 14.7 % ", (14.7, True)),
    ("0%", (0.0, True)),
    (None, (None, False)), ("n/a", (None, False)), ("%", (None, True)),
])
def test_leer_peso(valor, esperado):
    assert _leer_peso(valor) == esperado


@pytest.mark.parametrize("pesos,esperado", [
    ([0.6, 0.4, 0.0377], 100.0), ([84.19, 0.5], 1.0), ([0, None], 1.0), ([], 1.0),
])
def test_factor(pesos, esperado):
    assert _factor(pesos) == esperado


# ── vigencia.calcular_vigencia ──
def test_vigencia_anio_explicito_presente():
    metas = {2024: 10, 2025: 20, 2026: 30}
    assert calcular_vigencia(metas, 2025) == (2025, 2024, 20, 10)


def test_vigencia_anio_explicito_ausente_toma_anterior():
    metas = {2024: 10, 2025: 20, 2026: 30}
    # 2027 no existe → toma el <= más cercano (2026).
    assert calcular_vigencia(metas, 2027) == (2026, 2025, 30, 20)


def test_vigencia_sin_metas():
    assert calcular_vigencia({}, 2025) == (None, None, None, None)


# ── lector_fichas.codigo_de_hoja_ficha (6 convenciones de nombre) ──
@pytest.mark.parametrize("nombre,esperado", [
    ("Ficha técnica IR#1.1", "1.1"),
    ("Ficha técnica IP#1.1.1", "1.1.1"),
    ("R. 1.1", "1.1"),
    ("IR_1.1", "1.1"),
    ("1.1.1. Descripción del producto", "1.1.1"),
    # Juventud abrevia «Ficha» como «F» y cierra el código con punto.
    ("F IR#1.1.", "1.1"),
    ("F IP#7.4.5.", "7.4.5"),
    ("F IR#3.2", "3.2"),
    ("f ip#1.1.1", "1.1.1"),
    ("Plan de acción", None),
    ("Instructivo", None),
    # Una hoja que solo empieza por «F» no es una ficha.
    ("Formato 2024", None),
    ("Fuentes 1.2", None),
    ("F1.2", None),
])
def test_codigo_de_hoja_ficha(nombre, esperado):
    assert codigo_de_hoja_ficha(nombre) == esperado


def test_fichas_con_el_prefijo_abreviado_f_enriquecen_los_indicadores(tmp_path):
    """Plan con las hojas de ficha nombradas como en Juventud («F IR#1.1.», «F IP#1.1.1.»):
    `leer_fichas` las reconoce y entrega sus campos por código. Antes devolvía un
    diccionario vacío y el plan entero quedaba «sin ficha técnica»."""
    import openpyxl

    from extractor_pa.lector_fichas import leer_fichas

    wb = openpyxl.Workbook()
    wb.active.title = "Plan de acción"
    for nombre, metodologia in (("F IR#1.1.", "Encuesta anual"), ("F IP#1.1.1.", "Registro administrativo")):
        ws = wb.create_sheet(nombre)
        ws.cell(row=3, column=1, value="Metodología de medición")
        ws.cell(row=3, column=2, value=metodologia)
        ws.cell(row=4, column=1, value="Días de rezago")
        ws.cell(row=4, column=2, value="30")
    wb.create_sheet("Formato 2024").cell(row=3, column=1, value="Metodología de medición")
    ruta = tmp_path / "plan_fichas_f.xlsx"
    wb.save(ruta)

    fichas = leer_fichas(openpyxl.load_workbook(ruta))
    assert set(fichas) == {"1.1", "1.1.1"}
    assert fichas["1.1"] == {"metodologia": "Encuesta anual", "dias_rezago": 30}
    assert fichas["1.1.1"]["metodologia"] == "Registro administrativo"


# ── lector_fichas._leer_unidad (cuadrícula de opciones con la casilla a la derecha) ──
def _hoja_unidad(marcas=(), cual=None, celdas=()):
    """Hoja con el bloque «Unidad de medida» del formato real: tres opciones por fila
    (cols 3, 5 y 8), cada una con su casilla a la derecha (cols 4, 6 y 9)."""
    import openpyxl

    ws = openpyxl.Workbook().active
    ws.cell(row=6, column=1, value="Unidad de medida")
    for fila, (a, b, c) in enumerate((("Kilómetros", "kilos", "Tasa"),
                                      ("Hectáreas", "Metros", "Unidad productiva rural"),
                                      ("Personas", "Porcentaje", "otro")), start=8):
        ws.cell(row=fila, column=3, value=a)
        ws.cell(row=fila, column=5, value=b)
        ws.cell(row=fila, column=8, value=c)
    ws.cell(row=12, column=2, value="Cuál?")
    if cual:
        ws.cell(row=12, column=3, value=cual)
    ws.cell(row=14, column=1, value="Territorialización del indicador")
    ws.cell(row=16, column=4, value="X")            # una marca de OTRA sección: no cuenta
    for fila, columna in marcas:
        ws.cell(row=fila, column=columna, value="X")
    for fila, columna, valor in celdas:
        ws.cell(row=fila, column=columna, value=valor)
    return ws


@pytest.mark.parametrize("marcas,cual,esperado", [
    ([(10, 4)], None, "Personas"),                   # casilla de la primera columna
    # Hasta la 0.19.0 estas tres daban «Personas», None y None:
    ([(10, 6)], None, "Porcentaje"),                 # casilla de la segunda columna
    ([(8, 9)], None, "Tasa"),                        # casilla de la tercera columna
    ([(9, 9)], None, "Unidad productiva rural"),
    ([(8, 6)], None, "kilos"),
    # «otro» marcado: la unidad es lo escrito en «¿Cuál?».
    ([(10, 9)], "Puntaje", "Puntaje"),
    ([(10, 10)], "Componentes", "Componentes"),      # la «x» una celda más allá de «otro»
    ([(10, 9)], None, None),                         # «otro» marcado y sin respuesta
    # Sin ninguna marca: la respuesta a «¿Cuál?», como antes.
    ([], "Hogares", "Hogares"),
    ([], None, None),
    # Dos marcas: manda la primera en orden de lectura.
    ([(10, 6), (10, 9)], "Otra cosa", "Porcentaje"),
])
def test_leer_unidad_toma_la_opcion_pegada_a_la_casilla_marcada(marcas, cual, esperado):
    from extractor_pa.lector_fichas import _leer_unidad
    assert _leer_unidad(_hoja_unidad(marcas, cual)) == esperado


def test_leer_unidad_con_una_opcion_por_fila_y_la_casilla_a_la_izquierda():
    """Otras plantillas ponen la casilla ANTES de su opción («x  Otro  ¿Cuál?  Empresas»)
    o no tienen nada a la izquierda de la marca: se toma la opción más cercana."""
    import openpyxl

    from extractor_pa.lector_fichas import _leer_unidad

    ws = openpyxl.Workbook().active
    ws.cell(row=6, column=1, value="Unidad de medida")
    ws.cell(row=8, column=2, value="x")
    ws.cell(row=8, column=3, value="Hectáreas")
    assert _leer_unidad(ws) == "Hectáreas"

    ws = openpyxl.Workbook().active
    ws.cell(row=6, column=1, value="Unidad de medida")
    for columna, valor in ((2, "Otro"), (3, "X"), (4, "Cuál?"), (5, "Empresas")):
        ws.cell(row=8, column=columna, value=valor)
    assert _leer_unidad(ws) == "Empresas"


# ── pipeline._es_nombre_politica (C1: nombre vs decreto/CONPES/'No aplica') ──
@pytest.mark.parametrize("valor,esperado", [
    ("Política Pública de Servicios Públicos", True),
    ("POLÍTICA PÚBLICA DISTRITAL DE TRANSPARENCIA E INTEGRIDAD", True),
    ("233 de 2023", False),       # decreto
    ("01/2018", False),           # conpes/decreto con barra
    ("No aplica", False),
    ("", False), (None, False),
    ("Objetivo General de la Política Pública: lograr...", False),
    ("FORMATO DE PLAN DE ACCION POLÍTICAS PÚBLICAS", False),
])
def test_es_nombre_politica(valor, esperado):
    assert _es_nombre_politica(valor) is esperado
