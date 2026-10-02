# -*- coding: utf-8 -*-
"""
Plantilla antigua 2021–2025 (Vejez, Familias, Adultez, Negra-Afro v2…): los
encabezados van en las filas 10–12 y los datos desde la 13, no en 9–11 / 12.

Con las filas fijas el extractor encontraba los indicadores pero perdía las
metas anuales, la línea base, las fechas y todo el bloque financiero. Este plan
sintético reproduce esa estructura, incluido el título «Responsable de la
ejecución» combinado una columna antes de su bloque (como en Adultez v2).
"""

from __future__ import annotations

import os
import sys

from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from extractor_pa import MAPEO_NUEVO, extraer_plan_accion  # noqa: E402
from extractor_pa.pipeline import _ubicar_encabezados  # noqa: E402


def _plan_antiguo(ruta) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan de acción"
    ws.cell(row=4, column=2, value="Política Pública de Vejez de Prueba")
    celdas = {
        10: {1: "Objetivo específico", 2: "Importancia relativa  del objetivo específico (%)",
             3: "Indicadores de resultado", 13: "Indicadores de producto",
             22: "Meta de producto Final", 27: "Responsable de la ejecución"},
        11: {3: "Resultado esperado", 4: "Importancia relativa  del resultado (%)",
             5: "Nombre del indicador de resultado", 6: "Fórmula del indicador de resultado",
             7: "Enfoque", 8: "Tipo de anualización", 9: "Línea base",
             12: "Meta de resultado Final", 13: "Producto esperado",
             14: "Importancia relativa del producto (%)", 15: "Nombre indicador de producto",
             16: "Fórmula del indicador de producto", 17: "Enfoque", 18: "Tipo de anualización",
             19: "Línea base", 23: "2024", 27: "Costo total",
             28: "Sector ", 29: "Entidad", 30: "Dirección/Subdirección"},
        12: {9: "Valor", 10: "Año", 11: "Meta 2024", 19: "Valor", 20: "Año", 21: "Meta 2024",
             23: "Costo Estimado", 24: "Recurso disponible.", 25: "Fuente de financiación",
             26: "Código Proyecto de Inversión"},
        13: {1: "1. Objetivo uno", 2: 100, 3: "1.1 Resultado uno", 4: 100, 5: "IR uno",
             6: "a/b", 7: "Poblacional", 8: "Creciente", 9: 10, 10: 2020, 11: 20, 12: 20,
             13: "1.1.1 Producto uno", 14: 100, 15: "IP uno", 16: "c/d",
             17: "Género", 18: "Suma", 19: 0, 20: 2020, 21: 5, 22: 5,
             23: 1500, 24: 1200, 25: "Inversión", 26: "7872", 27: 1500,
             28: "Gestión Pública", 29: "Secretaría General", 30: "Alta Consejería"},
    }
    for fila, valores in celdas.items():
        for col, v in valores.items():
            ws.cell(row=fila, column=col, value=v)
    wb.save(ruta)


def test_ubicar_encabezados_por_resultado_esperado():
    wb = Workbook()
    ws = wb.active
    assert _ubicar_encabezados(ws, MAPEO_NUEVO) is MAPEO_NUEVO      # sin ancla: igual
    ws.cell(row=10, column=3, value="Resultado esperado")
    assert _ubicar_encabezados(ws, MAPEO_NUEVO) is MAPEO_NUEVO      # plantilla vigente
    ws.cell(row=10, column=3).value = None
    ws.cell(row=25, column=3, value="Resultado  esperado")
    m = _ubicar_encabezados(ws, MAPEO_NUEVO)
    assert (m.filas_encabezado, m.fila_datos) == ((24, 25, 26), 27)


def test_plan_antiguo_desplazado_completo(tmp_path):
    ruta = tmp_path / "plan_vejez_sintetico.xlsx"
    _plan_antiguo(ruta)
    res = extraer_plan_accion(ruta, anio_vigencia=2024)
    assert res.exitoso and res.metadatos.formato_detectado == "antiguo"

    [ir] = res.indicadores_resultado
    assert ir.metas_por_anio == {2024: 20}
    assert (ir.valor_linea_base, ir.anio_linea_base) == (10, 2020)
    assert (ir.tipo_anualizacion, ir.enfoque) == ("Creciente", "Poblacional")
    assert ir.meta_vigencia_actual == 20

    [ip] = res.indicadores_producto
    assert ip.metas_por_anio == {2024: 5}
    assert (ip.valor_linea_base, ip.tipo_anualizacion, ip.enfoque_principal) == (0, "Suma", "Género")
    assert ip.peso_pct == 100
    # El bloque de responsables se lee por sus subencabezados, no corrido una columna.
    assert (ip.sector_responsable, ip.entidad_responsable, ip.direccion_responsable) == (
        "Gestión Pública", "Secretaría General", "Alta Consejería")

    [fin] = res.financiero
    assert (fin.anio, fin.costo_estimado, fin.recurso_disponible) == (2024, 1500, 1200)
    assert (fin.fuente_financiacion, fin.codigo_proyecto) == ("Inversión", "7872")
